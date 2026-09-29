"""Every column and metric the Excel export writes explains itself (#7).

The workbook is meant to be handed to someone (or a chatbot) without the app
next to it, so each header carries a hover comment and a last "Definitions"
sheet lists them all. These tests fail when a new column lands without a
definition, which is the only way the two can drift apart.
"""

from __future__ import annotations

import inspect
import io
import re

import pytest
from openpyxl import load_workbook

from convexity import fetcher
from convexity import xlsx_export as xe


def _fetch_one_keys() -> set[str]:
    """Row keys fetch_one can put on a row, read from its source (and the
    helper that writes the currency-repaired multiples into the same dict)."""
    src = inspect.getsource(fetcher.fetch_one) + inspect.getsource(fetcher._currency_consistent)
    keys = set(re.findall(r'out\["([a-z0-9_]+)"\]', src))
    for n in (20, 50, 200):  # out[f"sma_{n}"] / out[f"above_sma_{n}"]
        keys |= {f"sma_{n}", f"above_sma_{n}"}
    return keys


def test_every_holdings_column_has_a_definition():
    cols = (
        xe.HOLDINGS_PRIMARY_COLS
        + xe.SENTIMENT_COLS
        + list(xe.EXTRA_LABELS.items())
        + xe.ANALYST_COLS
    )
    missing = [key for key, _ in cols if not (xe.COLUMN_DEFS.get(key) or ("", "", ""))[1]]
    assert not missing, f"COLUMN_DEFS has no meaning for: {missing}"


def test_every_portfolio_metric_has_a_definition():
    keys = [k for k, _ in xe.STATS_KEYS + xe.ANALYST_AGG_KEYS + xe.CONCENTRATION_KEYS]
    keys += [k for _, k in xe.OVERVIEW_METRICS]
    missing = [k for k in keys if not (xe.METRIC_DEFS.get(k) or ("", "", ""))[1]]
    assert not missing, f"METRIC_DEFS has no meaning for: {missing}"


def test_every_fetch_one_field_exports_with_a_label_and_a_definition():
    """A new fetch_one field lands in the holdings table as an 'extra'. It must
    get a readable header (EXTRA_LABELS) and a definition, or be skipped."""
    primary = {k for k, _ in xe.HOLDINGS_PRIMARY_COLS}
    extras = _fetch_one_keys() - primary - xe.HOLDINGS_SKIP_EXTRAS
    assert extras, "parsed no fetch_one keys: the regex no longer matches the source"
    unlabeled = sorted(k for k in extras if k not in xe.EXTRA_LABELS)
    undefined = sorted(k for k in extras if k not in xe.COLUMN_DEFS)
    assert not unlabeled, f"add an EXTRA_LABELS header (or skip) for: {unlabeled}"
    assert not undefined, f"add a COLUMN_DEFS entry for: {undefined}"


def _runner(rows, weight_sets, period, display_ccy):
    """A stand-in for analyze_portfolios_multi: fixed numbers, no network."""
    (set_name,) = weight_sets
    stats = {
        "total_return": 12.0,
        "ann_return": 12.0,
        "ann_vol": 18.0,
        "sharpe": 0.6,
        "max_dd": -9.0,
    }
    return {
        "results": {
            set_name: {
                "stats": stats,
                "benchmarks": {"SPY": {"stats": stats, "rel": {"beta": 1.1, "r2": 0.8}}},
                "analyst": {"mean_rating": 2.1, "n_analysts_total": 90, "holdings": []},
                "concentration": {"top5": 1.0, "herfindahl": 0.5, "effective_n": 2.0},
            }
        }
    }


@pytest.fixture
def workbook(monkeypatch):
    # The fallback analyst pull goes to Yahoo; the export must not need it here.
    monkeypatch.setattr(xe, "gather_analyst_info", lambda symbols, **kw: {})
    rows = [
        {
            "symbol": "AAPL",
            "name": "Apple",
            "price": 200.0,
            "market_cap": 3e12,
            "roe": 1.5,
            "sma_50": 190.0,
        },
        {
            "symbol": "MSFT",
            "name": "Microsoft",
            "price": 400.0,
            "market_cap": 3e12,
            "roe": 0.4,
            "sma_50": 380.0,
        },
    ]
    views = {"Demo": {"entries": "AAPL, MSFT", "rows": rows, "saved_at": "2026-09-29"}}
    blob = xe.build_workbook(views, analytics_runner=_runner, period="1Y")
    return load_workbook(io.BytesIO(blob))


def test_holdings_headers_carry_comments(workbook):
    ws = workbook["Demo"]
    header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Ticker")
    labels = {
        ws.cell(header_row, c).value: ws.cell(header_row, c) for c in range(1, ws.max_column + 1)
    }
    for label in ("Price", "P/E", "ROE", "SMA 50", "News read", "Target Upside (%)"):
        assert labels[label].comment is not None, label
    assert "Formula: net income (TTM) / shareholders' equity" in labels["ROE"].comment.text


def test_metric_labels_and_overview_headers_carry_comments(workbook):
    ws = workbook["Demo"]
    sharpe = next(
        ws.cell(r, 1) for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Sharpe Ratio"
    )
    assert sharpe.comment is not None
    assert "Formula: 252 x mean(r - rf)" in sharpe.comment.text
    hhi = next(
        ws.cell(r, 1) for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Herfindahl Index"
    )
    assert hhi.comment is not None
    ov = workbook["Overview"]
    assert ov.cell(4, 4).value == "1Y Ann Return %"
    assert ov.cell(4, 4).comment is not None
    assert ov.cell(4, 1).comment is None  # "Portfolio" needs no explanation


def test_definitions_sheet_is_last_and_complete(workbook):
    assert workbook.sheetnames[-1] == "Definitions"
    ws = workbook["Definitions"]
    header = [ws.cell(4, c).value for c in range(1, 6)]
    assert header == ["Section", "Metric", "Formula", "Meaning", "Typical range"]
    body = [[ws.cell(r, c).value for c in range(1, 6)] for r in range(5, ws.max_row + 1)]
    assert len(body) == len(xe.definition_rows())
    metrics = {row[1] for row in body}
    assert {"Sharpe Ratio", "Market read z (σ)", "Herfindahl Index", "SMA 200"} <= metrics
    assert all(row[3] for row in body), "every definition has a meaning"


def test_a_portfolio_named_definitions_does_not_collide(monkeypatch):
    monkeypatch.setattr(xe, "gather_analyst_info", lambda symbols, **kw: {})
    views = {"Definitions": {"entries": "AAPL", "rows": [{"symbol": "AAPL", "price": 1.0}]}}
    wb = load_workbook(io.BytesIO(xe.build_workbook(views)))
    assert wb.sheetnames == ["Overview", "Definitions", "Definitions-2"]

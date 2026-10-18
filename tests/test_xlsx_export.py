"""The Excel export: every column and metric explains itself, plus the sentiment columns.

The workbook is meant to be handed to someone (or a chatbot) without the app next to it,
so each header carries a hover comment and a last Definitions sheet lists them all.
These tests fail when a new column lands without a definition, which is the only way the
two can drift apart.

Also xlsx_export.SENTIMENT_COLS: the News read / Market read columns."""

import inspect
import io
import re

import pytest
import yfinance as yf
from openpyxl import load_workbook

from convexity import analytics, fetcher
from convexity import xlsx_export as xe  # noqa: E402


def _fetch_one_keys() -> set[str]:
    """Row keys fetch_one can put on a row, read from its source (and the
    helper that writes the currency-repaired multiples into the same dict)."""
    src = inspect.getsource(fetcher.fetch_one) + inspect.getsource(fetcher._currency_consistent)
    keys = set(re.findall(r'out\["([a-z0-9_]+)"\]', src))
    for n in (20, 50, 200):  # out[f"sma_{n}"] / out[f"above_sma_{n}"]
        keys |= {f"sma_{n}", f"above_sma_{n}"}
    keys |= {"target_low", "target_high", "target_median"}  # out[f"target_{k.lower()}"]
    return keys


def test_every_holdings_column_has_a_definition():
    cols = (
        xe.HOLDINGS_PRIMARY_COLS
        + xe.SENTIMENT_COLS
        + list(xe.EXTRA_LABELS.items())
        + xe.ANALYST_COLS
    )
    missing = [key for key, _ in cols if not (xe.COLUMN_DEFS.get(key) or ("", "", ""))[1]]
    assert not missing, f"COLUMN_DEFS has no meaning for: {missing}"


def test_every_portfolio_metric_has_a_definition():
    keys = [
        k for k, _ in xe.STATS_KEYS + xe.ANALYST_AGG_KEYS + xe.CONCENTRATION_KEYS + xe.WEIGHTED_KEYS
    ]
    keys += [k for _, k in xe.OVERVIEW_METRICS]
    missing = [k for k in keys if not (xe.METRIC_DEFS.get(k) or ("", "", ""))[1]]
    assert not missing, f"METRIC_DEFS has no meaning for: {missing}"


def test_every_fetch_one_field_exports_with_a_label_and_a_definition():
    """A new fetch_one field lands in the holdings table as an 'extra'. It must
    get a readable header (EXTRA_LABELS) and a definition, or be skipped."""
    primary = {k for k, _ in xe.HOLDINGS_PRIMARY_COLS}
    extras = _fetch_one_keys() - primary - xe.HOLDINGS_SKIP_EXTRAS
    assert extras, "parsed no fetch_one keys: the regex no longer matches the source"
    unlabeled = sorted(k for k in extras if k not in xe.EXTRA_LABELS)
    undefined = sorted(k for k in extras if k not in xe.COLUMN_DEFS)
    assert not unlabeled, f"add an EXTRA_LABELS header (or skip) for: {unlabeled}"
    assert not undefined, f"add a COLUMN_DEFS entry for: {undefined}"


def _runner(rows, weight_sets, period, display_ccy):
    """A stand-in for analyze_portfolios_multi: fixed numbers, no network."""
    (set_name,) = weight_sets
    stats = {
        "total_return": 12.0 if period == "1Y" else 80.0,
        "ann_return": 12.0,
        "ann_vol": 18.0,
        "sharpe": 0.6,
        "max_dd": -9.0,
    }
    spy = {**stats, "total_return": 10.0 if period == "1Y" else 70.0}
    return {
        "results": {
            set_name: {
                "period": period,
                "stats": stats,
                "weights_applied": {"AAPL": 0.5, "MSFT": 0.5},
                "benchmarks": {"SPY": {"stats": spy, "rel": {"beta": 1.1, "r2": 0.8}}},
                "weighted": {"pe": 30.0, "div_yield": 0.006, "coverage": {"pe": 1.0}},
                "analyst": {"mean_rating": 2.1, "n_analysts_total": 90, "holdings": []},
                "exposure": {"by_sector": {"Technology": 1.0}, "by_bucket": {"Mega ($200B+)": 1.0}},
                "concentration": {"top5": 1.0, "herfindahl": 0.5, "effective_n": 2.0},
                "contribution": [
                    {
                        "symbol": "AAPL",
                        "name": "Apple",
                        "weight": 0.5,
                        "period_return": 14.0,
                        "contribution": 7.0,
                        "sector": "Technology",
                    },
                    {
                        "symbol": "MSFT",
                        "name": "Microsoft",
                        "weight": 0.5,
                        "period_return": 10.0,
                        "contribution": 5.0,
                        "sector": "Technology",
                    },
                ],
                "warnings": ["MSFT (MSFT): before 1 Jan 2020 the prices are X's."],
            }
        }
    }


@pytest.fixture(autouse=True)
def no_yahoo(monkeypatch):
    """The export reads the analyst fields from the rows: never a Yahoo call."""

    def boom(*a, **k):
        raise AssertionError("the export called Yahoo")

    monkeypatch.setattr(yf, "Ticker", boom)
    monkeypatch.setattr(yf, "download", boom)


@pytest.fixture
def workbook(monkeypatch):
    rows = [
        {
            "symbol": "AAPL",
            "name": "Apple",
            "price": 200.0,
            "market_cap": 3e12,
            "roe": 1.5,
            "sma_50": 190.0,
        },
        {
            "symbol": "MSFT",
            "name": "Microsoft",
            "price": 400.0,
            "market_cap": 3e12,
            "roe": 0.4,
            "sma_50": 380.0,
        },
    ]
    views = {"Demo": {"entries": "AAPL, MSFT", "rows": rows, "saved_at": "2026-09-29"}}
    blob = xe.build_workbook(views, analytics_runner=_runner, period="1Y")
    return load_workbook(io.BytesIO(blob))


def test_holdings_headers_carry_comments(workbook):
    ws = workbook["Demo"]
    header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Ticker")
    labels = {
        ws.cell(header_row, c).value: ws.cell(header_row, c) for c in range(1, ws.max_column + 1)
    }
    for label in ("Price", "P/E", "ROE", "SMA 50", "News read", "Target Upside (%)"):
        assert labels[label].comment is not None, label
    assert "Formula: net income (TTM) / shareholders' equity" in labels["ROE"].comment.text


def test_metric_labels_and_overview_headers_carry_comments(workbook):
    ws = workbook["Demo"]
    sharpe = next(
        ws.cell(r, 1) for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Sharpe Ratio"
    )
    assert sharpe.comment is not None
    assert "Formula: 252 x mean(r - rf)" in sharpe.comment.text
    hhi = next(
        ws.cell(r, 1) for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Herfindahl Index"
    )
    assert hhi.comment is not None
    ov = workbook["Overview"]
    assert ov.cell(4, 4).value == "1Y Return"
    assert ov.cell(4, 4).comment is not None
    assert ov.cell(4, 1).comment is None  # "Portfolio" needs no explanation


def test_definitions_sheet_is_last_and_complete(workbook):
    assert workbook.sheetnames[-1] == "Definitions"
    ws = workbook["Definitions"]
    header = [ws.cell(4, c).value for c in range(1, 6)]
    assert header == ["Section", "Metric", "Formula", "Meaning", "Typical range"]
    body = [[ws.cell(r, c).value for c in range(1, 6)] for r in range(5, ws.max_row + 1)]
    assert len(body) == len(xe.definition_rows())
    metrics = {row[1] for row in body}
    assert {"Sharpe Ratio", "Market read z (σ)", "Herfindahl Index", "SMA 200"} <= metrics
    assert all(row[3] for row in body), "every definition has a meaning"


def test_a_portfolio_named_definitions_does_not_collide(monkeypatch):
    views = {"Definitions": {"entries": "AAPL", "rows": [{"symbol": "AAPL", "price": 1.0}]}}
    wb = load_workbook(io.BytesIO(xe.build_workbook(views)))
    assert wb.sheetnames == ["Read me", "Overview", "Definitions-2", "Data", "Definitions"]


def test_analyst_columns_fall_back_to_the_row(monkeypatch):
    """rec_key / n_analysts are not exported twice (HOLDINGS_SKIP_EXTRAS); the
    analyst columns show what the row carries, with no Yahoo call."""
    row = {
        "symbol": "AAPL",
        "price": 200.0,
        "rec_key": "buy",
        "n_analysts": 40,
        "recommendation_mean": 1.9,
        "target_low": 150.0,
        "target_mean_price": 250.0,
    }
    wb = load_workbook(io.BytesIO(xe.build_workbook({"Demo": {"entries": "AAPL", "rows": [row]}})))
    ws = wb["Demo"]
    hr = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Ticker")
    got = {ws.cell(hr, c).value: ws.cell(hr + 1, c).value for c in range(1, ws.max_column + 1)}
    assert got["Analyst Rating"] == "buy"
    assert got["# Analysts"] == 40
    assert got["Rec. Mean (1=Strong Buy, 5=Sell)"] == 1.9
    assert got["Target Low"] == 150.0 and got["Target Upside (%)"] == pytest.approx(25.0)
    assert "rec_key" not in got and "n_analysts" not in got and "target_low" not in got


def test_sheets_are_in_reading_order(workbook):
    assert workbook.sheetnames == ["Read me", "Overview", "Demo", "Data", "Definitions"]
    readme = [workbook["Read me"].cell(r, 2).value for r in range(1, 40)]
    assert any(
        v and v.startswith("Up 12.0% over the last year; the S&P 500 +10.0%.") for v in readme
    )


def test_portfolio_sheet_tells_the_story_first(workbook):
    ws = workbook["Demo"]
    assert ws["A3"].value.startswith("Up 12.0% over the last year")
    col_a = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    order = [
        "Performance",
        "Valuation",
        "Analysts",
        "Where the money is",
        "Concentration",
        "What drove the return",
        "Notes",
        "Holdings",
    ]
    found = [
        next(
            i
            for i, v in enumerate(col_a)
            if isinstance(v, str) and (v == t or v.startswith(t + " ("))
        )
        for t in order
    ]
    assert found == sorted(found)
    assert any(v and "the prices are X's" in v for v in col_a if isinstance(v, str))
    assert "Holdings1" in ws.tables
    perf = col_a.index("Performance") + 2  # 0-based list, header under the title
    header = [ws.cell(perf, c).value for c in range(1, 9)]
    assert header == [
        "Metric",
        "1Y Portfolio",
        "1Y S&P 500",
        "1Y Nasdaq-100",
        "5Y Portfolio",
        "5Y S&P 500",
        "5Y Nasdaq-100",
        "What it means",
    ]


def test_data_sheet_holds_every_figure_with_its_unit(workbook):
    ws = workbook["Data"]
    assert [ws.cell(4, c).value for c in range(1, 9)] == xe._Data.HEADER
    rows = {
        (ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 6).value): (
            ws.cell(r, 7).value,
            ws.cell(r, 8).value,
        )
        for r in range(5, ws.max_row + 1)
    }
    assert rows[("1Y", "total_return", "Portfolio")] == (12.0, "percent (12.5 = 12.5%)")
    assert rows[("5Y", "total_return", "S&P 500")][0] == 70.0
    assert rows[("1Y", "excess_vs_spy", "Portfolio")] == (2.0, "percentage points")
    assert rows[("1Y", "contribution.AAPL", "Portfolio")][0] == 7.0
    assert rows[("1Y", "by_sector.Technology", "Portfolio")] == (1.0, "fraction (0.125 = 12.5%)")


def test_a_stalled_run_does_not_hold_the_export(monkeypatch):
    monkeypatch.setattr(xe, "ANALYTICS_BUDGET_S", -1.0)
    views = {"Demo": {"entries": "AAPL", "rows": [{"symbol": "AAPL", "price": 1.0}]}}
    wb = load_workbook(io.BytesIO(xe.build_workbook(views, analytics_runner=_runner)))
    notes = [wb["Demo"].cell(r, 1).value or "" for r in range(1, wb["Demo"].max_row + 1)]
    assert any("not computed in time" in n for n in notes)


def test_prefetch_runs_once_over_every_portfolio():
    seen = []
    views = {
        n: {"entries": s, "rows": [{"symbol": s, "price": 1.0}]}
        for n, s in (("A", "AAPL"), ("B", "MSFT"))
    }
    xe.build_workbook(
        views,
        analytics_runner=_runner,
        prefetch=lambda rows, periods, ccy: seen.append(
            (sorted(r["symbol"] for r in rows), periods)
        ),
    )
    assert seen == [(["AAPL", "MSFT"], ["1Y", "5Y"])]


def test_prefetch_closes_downloads_the_longest_window_once(monkeypatch):
    import pandas as pd

    from convexity import cache

    calls = []
    idx = pd.bdate_range("2016-01-01", "2026-01-01")
    frame = pd.DataFrame({"ZZZ1": range(len(idx))}, index=idx, dtype=float)

    def fake(symbols, period, **kw):
        calls.append(period)
        cache._bulk_close_put("ZZZ1", period, frame["ZZZ1"])
        return frame

    monkeypatch.setattr(analytics, "_bulk_close", fake)
    monkeypatch.setattr(analytics, "_analyst_for", lambda s, r: {})
    monkeypatch.setattr("convexity.fx.prefetch_usd_series", lambda ccys, w: None)
    analytics.prefetch_closes([{"symbol": "ZZZ1"}], ["1Y", "5Y"])
    assert calls == ["10y"]
    two = cache._bulk_close_get_cached("ZZZ1", "2y")
    assert two.index[0] >= pd.Timestamp("2024-01-01") and two.index[-1] == idx[-1]


def test_no_target_range_call_for_a_holding_without_coverage(monkeypatch):
    def boom(symbol):
        raise AssertionError("asked Yahoo for a fund's price targets")

    monkeypatch.setattr(analytics, "_fetch_target_trio", boom)
    assert analytics._analyst_for("ZZZ2", {"price": 50.0})["target_low"] is None
    row = {"price": 50.0, "target_mean_price": 60.0, "target_low": 40.0,
           "target_high": 80.0, "target_median": 61.0}  # fmt: skip
    assert analytics._analyst_for("ZZZ3", row)["target_median"] == 61.0


S = {
    "news": {
        "tier": "bullish",
        "score": 0.9,
        "stale": False,
        "lenses": {
            "financials": {"score": 2.0},
            "outlook": {"score": -1.0},
            "competition": None,
            "regulation": None,
            "street": {"score": 1.0},
        },
    },
    "market": {"tier": "no_edge", "z": -0.42},
    "divergence": {
        "kind": "good_news_weak_reaction",
        "text": "Financials +2.0; Market read bearish",
    },
}


def test_every_sentiment_column_resolves():
    got = {label: xe._sentiment_value(key, S) for key, label in xe.SENTIMENT_COLS}
    assert got["News read"] == "bullish"
    assert got["News · Financials (−2..+2)"] == 2.0
    assert got["News · Outlook (−2..+2)"] == -1.0
    assert got["News · Competition (−2..+2)"] is None  # no news, not 0
    assert got["Market read"] == "no edge"
    assert got["Market read z (σ)"] == -0.42
    assert got["Divergence"].startswith("Financials +2.0")


def test_stale_and_missing_reads():
    stale = {"news": dict(S["news"], stale=True)}
    assert xe._sentiment_value("ns:news_tier", stale) == "bullish (stale)"
    assert xe._sentiment_value("ns:market_z", stale) is None
    assert all(xe._sentiment_value(k, None) is None for k, _ in xe.SENTIMENT_COLS)


def test_the_payload_blob_is_not_exported_raw():
    assert "news_sentiment" in xe.HOLDINGS_SKIP_EXTRAS

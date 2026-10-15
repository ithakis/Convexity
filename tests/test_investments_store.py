"""My Investments: the store, undo/redo, the page payload and the routes
(investments.py). Every network call is stubbed; public tickers only."""

import http.server
import importlib.util
import io
import json
import threading
import time
import urllib.error
import urllib.request
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from convexity import investments as inv
from convexity import paths, server

SEED = Path(__file__).resolve().parents[1] / "scripts" / "seed_demo_book.py"
QUOTE = {"AAPL": "USD", "MSFT": "USD", "NVDA": "USD", "SHEL.L": "GBp", "VWCE.DE": "EUR",
         "GOOGL": "USD"}  # fmt: skip
LAST = {"AAPL": 250.0, "MSFT": 500.0, "NVDA": 180.0, "SHEL.L": 2700.0, "VWCE.DE": 140.0,
        "GOOGL": 240.0}  # fmt: skip
FX = {"USD": 1.0, "GBP": 1.3, "EUR": 1.1}


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    """No network: quote currencies, names, splits, closes and FX are fixed."""
    inv._SPLITS.clear()
    inv._SPLITS_GOOD.clear()
    inv._PAGE_MEMO.clear()
    inv._QCCY.clear()
    monkeypatch.setattr(inv, "_fetch_quote_ccy", lambda s: QUOTE.get(s))
    monkeypatch.setattr(inv, "_pack_hit", lambda t: None)
    monkeypatch.setattr(inv, "_resolve", lambda text: text.strip().upper())
    monkeypatch.setattr(
        inv, "_fetch_splits", lambda s: [("2024-06-10", 10.0)] if s == "NVDA" else []
    )
    monkeypatch.setattr(inv, "_fx_now", lambda ccy, base: FX[ccy] / FX[base])

    def closes(symbols, period):
        # The last close on the last weekday before today: the book's calendar
        # has weekdays only, so a Saturday close would never be read.
        last = date.today() - timedelta(days=1)
        while last.weekday() >= 5:
            last -= timedelta(days=1)
        days = pd.to_datetime([last - timedelta(days=45), last])
        return pd.DataFrame({s: [LAST[s] * 0.9, LAST[s]] for s in symbols if s in LAST}, index=days)

    monkeypatch.setattr(inv, "_closes", closes)
    monkeypatch.setattr(
        inv,
        "_bench_closes",
        lambda period, ticker: pd.Series([400.0], index=pd.to_datetime(["2000-01-03"])),
    )
    monkeypatch.setattr(inv, "_dividends", lambda s: None)
    monkeypatch.setattr(inv, "_rf_closes", lambda period: None)
    monkeypatch.setattr(
        inv,
        "_fx_mult",
        lambda ccys, disp, dates: {c: np.full(len(dates), FX[c] / FX[disp]) for c in ccys},
    )


def _seed():
    spec = importlib.util.spec_from_file_location("seed_demo_book", SEED)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    path = paths.state_file("investments")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(mod.build_book()), encoding="utf-8")


def _raw():
    return json.loads(paths.state_file("investments").read_text(encoding="utf-8"))


BUY = {"type": "buy", "symbol": "AAPL", "date": "2025-01-02", "qty": 10, "price": 200, "fee": 1}


# ------------------------------------------------------------------- store


def test_no_book_is_empty():
    assert inv.book() == {"empty": True, "rev": 0, "undo_label": None, "redo_label": None}
    assert inv.book_symbols() == []


def test_add_update_delete_undo_redo():
    b = inv.add(BUY, 0)
    assert b["rev"] == 1 and b["positions"][0]["shares"] == 10
    assert b["undo_label"] == "Added buy · 10 AAPL" and b["redo_label"] is None
    eid = b["entries"][0]["id"]
    b = inv.update(eid, {"qty": 12}, 1)
    assert b["positions"][0]["shares"] == 12 and b["undo_label"].startswith("Edited buy")
    b = inv.undo(2)
    assert b["positions"][0]["shares"] == 10 and b["redo_label"].startswith("Edited")
    b = inv.redo(3)
    assert b["positions"][0]["shares"] == 12
    b = inv.delete(eid, 4)
    assert b["empty"] is True and b["undo_label"].startswith("Deleted buy")
    b = inv.undo(5)
    assert b["positions"][0]["shares"] == 12
    # A new change after an undo clears redo.
    inv.undo(6)
    b = inv.add(BUY | {"symbol": "MSFT"}, 7)
    assert b["redo_label"] is None
    assert len(_raw()["redo"]) == 0


def test_undo_restores_an_entry_at_its_old_position():
    inv.add(BUY, 0)
    inv.add(BUY | {"qty": 5}, 1)
    first = _raw()["entries"][0]["id"]
    inv.delete(first, 2)
    inv.undo(3)
    assert _raw()["entries"][0]["id"] == first


def test_stale_rev_is_refused():
    inv.add(BUY, 0)
    with pytest.raises(inv.StaleRev):
        inv.add(BUY, 0)
    with pytest.raises(inv.StaleRev):
        inv.undo("not a number")
    assert _raw()["rev"] == 1


def test_journal_is_capped():
    rev = 0
    for _ in range(inv.JOURNAL_MAX + 5):
        rev = inv.add(BUY | {"qty": 1}, rev)["rev"]
    assert len(_raw()["journal"]) == inv.JOURNAL_MAX


def test_oversell_is_refused_in_plain_words_and_nothing_is_saved():
    inv.add(BUY, 0)
    with pytest.raises(inv.Refused) as ex:
        inv.add(
            {"type": "sell", "symbol": "AAPL", "date": "2025-02-03", "qty": 11, "price": 210}, 1
        )
    assert ex.value.status == 422
    assert ex.value.message == "That sells 11 AAPL on 3 Feb 2025, but you hold 10 then."
    assert _raw()["rev"] == 1 and len(_raw()["entries"]) == 1


def test_an_edit_that_breaks_a_later_sale_names_that_sale():
    b = inv.add(BUY, 0)
    eid = b["entries"][0]["id"]
    inv.add({"type": "sell", "symbol": "AAPL", "date": "2025-03-03", "qty": 10, "price": 210}, 1)
    with pytest.raises(inv.Refused) as ex:
        inv.update(eid, {"qty": 5}, 2)
    assert "sale of 10" in ex.value.message and "3 Mar 2025" in ex.value.message


@pytest.mark.parametrize(
    "fields, words",
    [
        (
            {"type": "buy", "symbol": "AAPL", "date": "2025-01-02", "qty": 0, "price": 1},
            "more than zero",
        ),
        ({"type": "buy", "symbol": "AAPL", "date": "2999-01-02", "qty": 1, "price": 1}, "future"),
        (
            {"type": "buy", "symbol": "", "date": "2025-01-02", "qty": 1, "price": 1},
            "Choose a company",
        ),
        ({"type": "deposit", "date": "2025-01-02", "amount": "abc"}, "number"),
        (
            {"type": "dividend", "symbol": "AAPL", "date": "2025-01-02", "amount": 5, "tax": 6},
            "tax",
        ),
        ({"type": "split", "symbol": "AAPL", "date": "2025-01-02", "ratio": 1}, "ratio"),
        ({"type": "nonsense", "date": "2025-01-02"}, "kind of entry"),
        ({"type": "buy", "symbol": "ZZZZ", "date": "2025-01-02", "qty": 1, "price": 1}, "Yahoo"),
    ],
)
def test_bad_fields_are_refused(fields, words):
    with pytest.raises(inv.BookError) as ex:
        inv.add(fields, 0)
    assert words in ex.value.message
    assert not paths.state_file("investments").exists()


def test_a_malformed_book_is_never_overwritten():
    path = paths.state_file("investments")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("{not json", encoding="utf-8")
    b = inv.book()
    assert b["empty"] is True and "damaged" in b["error"]
    with pytest.raises(inv.BookUnreadable) as ex:
        inv.add(BUY, 0)
    assert ex.value.status == 409
    assert path.read_text(encoding="utf-8") == "{not json"
    assert inv.book_symbols() == []


def test_a_newer_schema_is_refused_not_rewritten():
    path = paths.state_file("investments")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"version": 99, "entries": []}), encoding="utf-8")
    with pytest.raises(inv.BookUnreadable):
        inv.add(BUY, 0)


def test_every_write_keeps_a_backup_of_the_previous_file():
    inv.add(BUY, 0)
    inv.add(BUY | {"qty": 3}, 1)
    bak = json.loads(paths.state_file("investments").with_name("investments.json.bak").read_text())
    assert bak["rev"] == 1


def test_logs_never_carry_amounts(capsys):
    b = inv.add(BUY | {"qty": 123457, "price": 98765.4}, 0)
    inv.update(b["entries"][0]["id"], {"price": 55555}, 1)
    out = capsys.readouterr().out
    assert "[investments] add" in out and "AAPL" in out
    for secret in ("123457", "98765", "55555"):
        assert secret not in out


@pytest.mark.parametrize(
    "text",
    [
        json.dumps({"version": "x", "entries": []}),
        json.dumps({"rev": "x", "entries": []}),
        json.dumps({"entries": [1]}),
        json.dumps({"entries": [{"type": "buy"}]}),
        json.dumps({"entries": [{"type": "buy", "date": "2025-01-02"}]}),
    ],
)
def test_a_book_with_a_bad_shape_is_unreadable_not_a_crash(text):
    path = paths.state_file("investments")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    assert "error" in inv.book() and inv.book_symbols() == []
    with pytest.raises(inv.BookUnreadable):
        inv.add(BUY, 0)
    assert path.read_text(encoding="utf-8") == text


@pytest.mark.parametrize(
    "patch, words",
    [
        ({"qty": True}, "number"),
        ({"price": 1e300}, "too large"),
        ({"ccy": "<b>x</b>"}, "currency"),
        ({"ccy": ["USD"]}, "currency"),
        ({"symbol": "A" * 500}, "company"),
    ],
)
def test_odd_values_are_refused(patch, words):
    with pytest.raises(inv.BookError) as ex:
        inv.add(BUY | patch, 0)
    assert words in ex.value.message


def test_only_known_sources_are_stored():
    inv.add(BUY | {"source": {"x": 1}}, 0)
    inv.add(BUY | {"source": "quick"}, 1)
    assert [e["source"] for e in _raw()["entries"]] == ["manual", "quick"]


def test_changing_the_company_of_an_entry_takes_the_new_currency():
    b = inv.add(BUY | {"symbol": "SHEL.L", "fx_rate": 1.25}, 0)
    eid = b["entries"][0]["id"]
    inv.update(eid, {"symbol": "VWCE.DE"}, 1)
    e = _raw()["entries"][0]
    assert e["symbol"] == "VWCE.DE" and e["ccy"] == "EUR" and e["fx_rate"] is None


def test_a_pence_holding_without_a_yahoo_price_is_shown_in_pence(monkeypatch):
    monkeypatch.setattr(inv, "_closes", lambda symbols, period: pd.DataFrame())
    b = inv.add(BUY | {"symbol": "SHEL.L", "price": 25.0, "fee": 0}, 0)
    p = b["positions"][0]
    assert p["quote_ccy"] == "GBp" and p["price"] == pytest.approx(2500)
    assert p["avg_price"] == pytest.approx(2500) and p["note"]


# ----------------------------------------------------------------- payload


def test_seeded_demo_book_replays_through_the_engine():
    _seed()
    b = inv.book()
    assert b["demo"] is True and b["empty"] is False and not b["warnings"]
    pos = {p["symbol"]: p for p in b["positions"]}
    assert pos["NVDA"]["shares"] == 120  # 20 pre-split x 10, less 80 sold
    assert pos["MSFT"]["shares"] == 50 and pos["SHEL.L"]["shares"] == 300
    assert pos["SHEL.L"]["avg_price"] == pytest.approx(2541.0)  # pence, fee included
    h = b["headline"]
    # Value = positions + cash; weights + cash = 100%; profit = value - money in.
    assert sum(p["value"] for p in b["positions"]) + h["cash"] == pytest.approx(h["value"])
    assert sum(p["weight"] for p in b["positions"]) + h["cash_weight"] == pytest.approx(100)
    assert h["profit_total"] == pytest.approx(h["value"] - h["net_deposits"])
    assert h["profit_held"] + h["profit_sold"] + h["dividends"] - h["fees"] == pytest.approx(
        h["profit_total"]
    )
    assert h["net_deposits"] == pytest.approx(65000)
    # Performance (Phase 3): the chart's last point is the headline value.
    s = b["series"]
    assert s["value"][-1] == pytest.approx(h["value"], abs=0.01)
    assert s["invested"][-1] == pytest.approx(h["net_deposits"], abs=0.01)
    assert all(isinstance(h[k][p], float) for k in ("twr", "bench_twr", "mwr") for p in inv.PERIODS)
    # FX split only for holdings quoted in another currency.
    assert pos["MSFT"]["profit_fx"] is None
    assert pos["VWCE.DE"]["profit_price"] + pos["VWCE.DE"]["profit_fx"] == pytest.approx(
        pos["VWCE.DE"]["profit_held"]
    )
    assert inv.book_symbols() == sorted(pos)


def test_a_real_book_is_computed_and_never_demo():
    b = inv.add(BUY, 0)
    assert b["demo"] is False
    p = b["positions"][0]
    assert p["value"] == pytest.approx(2500) and p["profit_held"] == pytest.approx(2500 - 2001)
    assert b["headline"]["implied_deposits"] == pytest.approx(2001)


def test_period_starts_clamp_month_ends_and_use_jan_1_for_ytd():
    dates = ["2025-12-31", "2026-01-02", "2026-02-27", "2026-03-02", "2026-03-31"]
    st = inv.period_starts(dates)
    assert dates[st["1M"]] == "2026-02-27"
    assert dates[st["YTD"]] == "2025-12-31"
    assert st["ALL"] == 0


# ------------------------------------------------------------------ routes


@pytest.fixture
def srv():
    httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    threading.Thread(
        target=httpd.serve_forever, kwargs={"poll_interval": 0.05}, daemon=True
    ).start()
    yield f"http://127.0.0.1:{httpd.server_address[1]}"
    httpd.shutdown()
    httpd.server_close()


def _post(base, path, body, ctype="application/json"):
    req = urllib.request.Request(base + path, method="POST", data=json.dumps(body).encode(),
                                 headers={"Content-Type": ctype})  # fmt: skip
    try:
        with urllib.request.urlopen(req, timeout=15) as r:
            return r.status, json.loads(r.read())
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read())


def _delete(base, path):
    req = urllib.request.Request(base + path, method="DELETE")
    try:
        with urllib.request.urlopen(req, timeout=15) as r:
            return r.status
    except urllib.error.HTTPError as e:
        return e.code


def test_routes_round_trip(srv):
    code, b = _post(srv, "/api/investments/entries", {"base_rev": 0, "entry": BUY})
    assert code == 200 and b["rev"] == 1
    code, b2 = _post(srv, "/api/investments/entries", {"base_rev": 0, "entry": BUY})
    assert code == 409 and b2["error"] == "StaleRev"
    code, b3 = _post(
        srv,
        "/api/investments/entries",
        {"base_rev": 1, "entry": BUY | {"type": "sell", "qty": 99, "date": "2025-02-03"}},
    )
    assert code == 422 and "you hold 10" in b3["message"]
    code, b4 = _post(srv, "/api/investments/undo", {"base_rev": 1})
    assert code == 200 and b4["empty"] is True
    code, b5 = _post(srv, "/api/investments/redo", {"base_rev": 2})
    assert code == 200 and b5["positions"][0]["symbol"] == "AAPL"
    with urllib.request.urlopen(srv + "/api/investments/symbols") as r:
        assert json.loads(r.read()) == {"symbols": ["AAPL"]}
    assert _post(srv, "/api/investments/undo", {"base_rev": 3}, "text/plain")[0] == 403


@pytest.mark.parametrize(
    "path, body",
    [
        ("/api/watchlists", {"name": "__book__", "entries": "AAPL"}),
        ("/api/portfolio/rename", {"old": "__book__", "new": "Mine"}),
        ("/api/portfolio/rename", {"old": "Mine", "new": "__book__"}),
        ("/api/mpt-runs", {"view": "__book__", "run": {}}),
        ("/api/weight-presets", {"view": "__book__", "name": "p", "weights": {"AAPL": 1}}),
        ("/api/weight-presets/active", {"view": "__book__", "name": None}),
    ],
)
def test_the_book_tab_name_is_reserved(srv, path, body):
    code, out = _post(srv, path, body)
    assert code == 400 and "read-only" in out["error"]


def test_the_book_tab_cannot_be_deleted_but_its_rows_can_be_cached(srv):
    assert _delete(srv, "/api/views/__book__") == 400
    assert _delete(srv, "/api/weight-presets?view=__book__&name=p") == 400
    code, _ = _post(srv, "/api/views/__book__", {"entries": "AAPL", "rows": [{"symbol": "AAPL"}]})
    assert code == 200


def test_seed_refuses_without_convexity_home(monkeypatch):
    import subprocess
    import sys

    monkeypatch.delenv("CONVEXITY_HOME", raising=False)
    r = subprocess.run([sys.executable, str(SEED)], capture_output=True, text=True)
    assert r.returncode == 2
    assert "Refusing" in r.stdout


@pytest.mark.parametrize("entry", [[1, 2], "buy", 5])
def test_an_entry_that_is_not_an_object_is_a_400(srv, entry):
    code, out = _post(srv, "/api/investments/entries", {"base_rev": 0, "entry": entry})
    assert code == 400


def test_a_preset_may_be_called_like_the_book_tab(srv):
    _post(srv, "/api/views/Mine", {"entries": "AAPL", "rows": [{"symbol": "AAPL"}]})
    body = {"view": "Mine", "name": "__book__", "weights": {"AAPL": 1}}
    code, out = _post(srv, "/api/weight-presets", body)
    assert code == 200, out


# ------------------------------------------------------- import (Phases 5-6)


def _batch():
    return [
        {"type": "deposit", "date": "2024-03-01", "amount": 5000, "ccy": "USD"},
        {
            "type": "buy",
            "date": "2024-03-12",
            "symbol": "AAPL",
            "qty": 10,
            "price": 182.4,
            "qty_basis": "trade",
        },
        {
            "type": "buy",
            "date": "2024-01-02",
            "symbol": "MSFT",
            "qty": 5,
            "price": 370.0,
            "qty_basis": "current",
        },
    ]


def test_an_import_is_one_batch_and_one_undo():
    b = inv.add_batch(_batch(), 0)
    assert b["undo_label"] == "Imported 3 entries" and len(b["entries"]) == 3
    raw = _raw()
    assert {e["source"] for e in raw["entries"]} == {"import"}
    assert len({e["batch"] for e in raw["entries"]}) == 1 and len(raw["journal"]) == 1
    b = inv.undo(b["rev"])
    assert b["empty"] is True
    assert len(inv.redo(b["rev"])["entries"]) == 3


def test_a_bad_row_names_itself_and_nothing_is_saved():
    rows = _batch()
    rows[1] = rows[1] | {"qty": -3}
    with pytest.raises(inv.BookError) as ex:
        inv.add_batch(rows, 0)
    assert ex.value.message.startswith("Row 2:")
    assert inv.book()["empty"] is True


def test_the_dry_run_flags_a_sale_in_the_statements_share_count():
    # 5 NVIDIA sold in Feb 2024 were 50 shares after the 10:1 split; the
    # owner's statement says 5, so the words (and the shortfall) say 5.
    probs = inv.check_rows([
        {"type": "sell", "date": "2024-02-21", "symbol": "NVDA", "qty": 5, "price": 674.1, "qty_basis": "trade"},
        {"type": "buy", "date": "2024-03-12", "symbol": "AAPL", "qty": 10, "price": 182.4, "qty_basis": "trade"},
    ])  # fmt: skip
    assert list(probs) == [0] and probs[0]["kind"] == "oversell"
    assert probs[0]["short"] == pytest.approx(5) and "Sells 5 " in probs[0]["message"]
    assert inv.book()["empty"] is True  # a dry run writes nothing


def test_corrected_names_are_remembered_only_for_booked_tickers():
    inv.add_batch(_batch(), 0, aliases={"Apple ": "AAPL", "Tesla": "TSLA"})
    assert _raw()["settings"]["aliases"] == {"apple": "AAPL"}
    ctx = inv.import_context()
    assert ctx["aliases"] == {"apple": "AAPL"}
    assert ctx["held"] == {"AAPL": pytest.approx(10), "MSFT": pytest.approx(5)}


def test_import_routes(srv, monkeypatch):
    from convexity import importer
    from convexity import news_sentiment as ns

    monkeypatch.setattr(ns, "NVIDIA_API_KEY", "")
    code, b = _post(srv, "/api/investments/import", {"text": "10 Apple"})
    assert code == 409 and "NVIDIA key" in b["message"]
    monkeypatch.setattr(ns, "NVIDIA_API_KEY", "test-key")
    monkeypatch.setattr(
        importer, "_structure", lambda text, today, base: {"items": [], "skipped": []}
    )
    code, b = _post(srv, "/api/investments/import", {"text": "nothing here"})
    assert code == 200 and b["items"] == []
    code, b = _post(
        srv, "/api/investments/import/check", {"entries": [_batch()[1] | {"type": "sell"}]}
    )
    assert code == 200 and b["problems"]["0"]["kind"] == "oversell"
    code, b = _post(srv, "/api/investments/import/apply", {"base_rev": 0, "entries": _batch()})
    assert code == 200 and b["undo_label"] == "Imported 3 entries"
    code, _ = _post(
        srv, "/api/investments/import/apply", {"base_rev": 0, "entries": _batch()}, "text/plain"
    )
    assert code == 403


def test_the_dry_run_refuses_bad_shapes_in_plain_words():
    # Review finding: a non-object row was a 500.
    assert inv.check_rows(["x", _batch()[0]]) == {
        0: {"kind": "invalid", "message": "This row couldn't be read."}
    }
    with pytest.raises(inv.BookError):
        inv.check_rows({"not": "a list"})


def test_routes_answer_in_the_display_currency(srv):
    code, b = _post(srv, "/api/investments/entries", {"base_rev": 0, "entry": BUY, "ccy": "GBP"})
    assert code == 200 and b["ccy"] == "GBP"
    assert b["headline"]["value"] == pytest.approx(2500 / 1.3)
    with urllib.request.urlopen(srv + "/api/investments?ccy=EUR") as r:
        assert json.loads(r.read())["ccy"] == "EUR"
    with urllib.request.urlopen(srv + "/api/investments?ccy=nope") as r:
        assert json.loads(r.read())["ccy"] == "USD"
    code, b = _post(srv, "/api/investments/undo", {"base_rev": 1, "ccy": ["GBP"]})
    assert code == 200 and b.get("ccy") in (None, "USD")  # a non-string ccy is ignored


def test_book_refresh_route_needs_a_book(srv, monkeypatch):
    from convexity import jobs

    code, b = _post(srv, "/api/refresh-job", {"view": jobs.BOOK_VIEW, "phases": ["book"]})
    assert code == 400 and b["error"] == "nothing to refresh"
    monkeypatch.setattr(inv, "refresh_market", lambda symbols, **kw: None)
    inv.add(BUY, 0)
    code, b = _post(srv, "/api/refresh-job", {"view": jobs.BOOK_VIEW, "phases": ["book"]})
    assert code == 202 and b["job_id"]
    job = jobs.get(b["job_id"])
    for _ in range(200):
        if job.is_terminal():
            break
        threading.Event().wait(0.02)
    assert job.state == "done" and job.views == [jobs.BOOK_VIEW] and job.phases == ("book",)


# ------------------------------------------------------------------- clear


def test_clear_is_one_undoable_step_with_a_copy_kept():
    inv.add(BUY, 0)
    inv.add(BUY | {"symbol": "MSFT", "price": 400}, 1)
    raw = _raw()
    raw["settings"]["aliases"] = {"apple": "AAPL"}
    paths.state_file("investments").write_text(json.dumps(raw), encoding="utf-8")
    b = inv.clear(2)
    assert b["empty"] is True and b["undo_label"] == "Cleared 2 entries"
    assert _raw()["settings"]["aliases"] == {"apple": "AAPL"}
    copies = list(paths.state_file("investments").parent.glob("investments.before-clear-*.json"))
    assert len(copies) == 1 and len(json.loads(copies[0].read_text())["entries"]) == 2
    inv.undo(3)
    assert _raw()["entries"] == raw["entries"]  # every entry back, in its place
    with pytest.raises(inv.StaleRev):
        inv.clear(3)


def test_clear_keeps_the_newest_three_copies(monkeypatch):
    inv.add(BUY, 0)
    folder = paths.state_file("investments").parent
    for k in range(4):
        (folder / f"investments.before-clear-2026010{k}-000000.json").write_text("{}")
    inv.clear(1)
    left = sorted(p.name for p in folder.glob("investments.before-clear-*.json"))
    assert len(left) == 3 and "20260100" not in " ".join(left)


def test_nothing_to_clear_is_said_plainly():
    with pytest.raises(inv.BookError, match="nothing to clear"):
        inv.clear(0)


def test_apply_keeps_the_opening_flag_of_imported_holdings():
    """The review sends ``opening`` for every holding row; Apply used to drop
    it (2026-10-11), booking each holding as bought on its date."""
    rows = [{"type": "buy", "symbol": "AAPL", "date": "2025-01-02", "qty": 10, "price": 150,
             "qty_basis": "current", "opening": True},
            {"type": "buy", "symbol": "MSFT", "date": "2025-01-02", "qty": 1, "price": 300,
             "qty_basis": "trade"}]  # fmt: skip
    inv.add_batch(rows, 0)
    saved = {e["symbol"]: e for e in _raw()["entries"]}
    assert saved["AAPL"]["opening"] is True
    assert not saved["MSFT"].get("opening")


# ------------------------------------------------- split history (Phase 4)
# A failed split fetch is not "no splits": the seeded book's 20 NVIDIA bought
# before the 10:1 split and the 80 sold after it oversell at face value, so a
# Yahoo outage used to show a sound book as broken.

NVDA_SPLITS = [("2024-06-10", 10.0)]
_REAL_FETCH_SPLITS = inv._fetch_splits  # read at import, before the offline stub


def _yahoo_down_for(monkeypatch, *down):
    monkeypatch.setattr(
        inv, "_fetch_splits", lambda s: None if s in down else (NVDA_SPLITS if s == "NVDA" else [])
    )


def test_a_failed_split_fetch_leaves_the_holding_out_never_broken(monkeypatch):
    _seed()
    _yahoo_down_for(monkeypatch, "NVDA")
    b = inv.book()
    assert "broken" not in b and b["headline"] is not None
    assert [p["symbol"] for p in b["splits_pending"]] == ["NVDA"]
    nv = [p for p in b["positions"] if p["symbol"] == "NVDA"]
    assert nv == [b["splits_pending"][0]] and nv[0]["pending"] is True
    assert b["positions"][-1]["symbol"] == "NVDA"  # last, with no figures
    assert any("left out of the totals" in w for w in b["warnings"])
    # The rest is the book without NVIDIA, computed in full.
    h, held = b["headline"], [p for p in b["positions"] if not p.get("pending")]
    assert sum(p["value"] for p in held) + h["cash"] == pytest.approx(h["value"])
    assert h["profit_total"] == pytest.approx(h["value"] - h["net_deposits"])


def test_a_failure_after_a_good_answer_keeps_the_last_good_list(monkeypatch):
    _seed()
    good = inv.book()
    clock = [1e9]
    monkeypatch.setattr(inv.time, "time", lambda: clock[0])
    inv._SPLITS.clear()
    inv.book()  # fetched at the fake clock
    _yahoo_down_for(monkeypatch, "NVDA")
    clock[0] += inv._SPLITS_TTL + 1  # due again, and Yahoo fails
    b = inv.book()
    assert not b["splits_pending"] and b["headline"]["value"] == pytest.approx(
        good["headline"]["value"]
    )
    assert inv._SPLITS["NVDA"][1] is False  # failure noted: retried after _MISS_TTL
    calls = []
    monkeypatch.setattr(inv, "_fetch_splits", lambda s: calls.append(s) or NVDA_SPLITS)
    clock[0] += inv._MISS_TTL - 1
    inv.market_splits(["NVDA"])
    assert calls == []
    clock[0] += 2
    inv.market_splits(["NVDA"])
    assert calls == ["NVDA"] and inv._SPLITS["NVDA"][1] is True


def test_good_split_lists_survive_a_restart_and_a_bad_file_is_ignored(monkeypatch):
    _seed()
    inv.book()
    saved = json.loads(paths.state_file("splits").read_text(encoding="utf-8"))
    assert saved["NVDA"] == [["2024-06-10", 10.0]]
    # A restart (memory gone) during an outage still has the list.
    inv._SPLITS.clear()
    inv._SPLITS_GOOD.clear()
    monkeypatch.setattr(inv, "_splits_loaded", None)
    _yahoo_down_for(monkeypatch, "NVDA", "MSFT", "AAPL", "SHEL.L", "VWCE.DE")
    assert not inv.book()["splits_pending"]
    paths.state_file("splits").write_text("{not json", encoding="utf-8")
    inv._SPLITS.clear()
    inv._SPLITS_GOOD.clear()
    monkeypatch.setattr(inv, "_splits_loaded", None)
    assert [p["symbol"] for p in inv.book()["splits_pending"]] == sorted(QUOTE.keys() - {"GOOGL"})


def test_fetch_splits_tells_a_failure_from_no_splits(monkeypatch):
    import yfinance as yf

    frames = {
        "DOWN": pd.DataFrame(),
        "NONE": pd.DataFrame({"Close": [1.0, 2.0], "Stock Splits": [0.0, 0.0]},
                             index=pd.to_datetime(["2024-01-02", "2024-01-03"])),
        "TEN": pd.DataFrame({"Close": [1.0, 2.0], "Stock Splits": [0.0, 10.0]},
                            index=pd.to_datetime(["2024-06-07", "2024-06-10"])),
    }  # fmt: skip

    class T:
        def __init__(self, s):
            self.s = s

        def history(self, **kw):
            if self.s == "RAISE":
                raise RuntimeError("rate limited")
            return frames[self.s]

    monkeypatch.setattr(yf, "Ticker", T)
    assert _REAL_FETCH_SPLITS("DOWN") is None and _REAL_FETCH_SPLITS("RAISE") is None
    assert _REAL_FETCH_SPLITS("NONE") == []
    assert _REAL_FETCH_SPLITS("TEN") == [("2024-06-10", 10.0)]


def test_an_edit_to_a_pending_holding_waits_and_others_go_through(monkeypatch):
    _seed()
    rev = inv.book()["rev"]
    inv._SPLITS.clear()
    inv._SPLITS_GOOD.clear()
    _yahoo_down_for(monkeypatch, "NVDA")
    with pytest.raises(inv.Refused) as ex:
        inv.add({"type": "sell", "symbol": "NVDA", "date": "2025-03-03", "qty": 10, "price": 110,
                 "qty_basis": "trade"}, rev)  # fmt: skip
    assert "split history for NVIDIA didn't load" in ex.value.message
    assert _raw()["rev"] == rev  # nothing saved
    b = inv.add(BUY, rev)
    assert b["rev"] == rev + 1 and [p["symbol"] for p in b["splits_pending"]] == ["NVDA"]


def test_the_dry_run_says_a_pending_row_cant_be_checked_yet(monkeypatch):
    _yahoo_down_for(monkeypatch, "NVDA")
    probs = inv.check_rows([
        {"type": "sell", "date": "2024-02-21", "symbol": "NVDA", "qty": 5, "price": 674.1, "qty_basis": "trade"},
        {"type": "buy", "date": "2024-03-12", "symbol": "AAPL", "qty": 10, "price": 182.4, "qty_basis": "trade"},
    ])  # fmt: skip
    assert list(probs) == [0] and probs[0]["kind"] == "splits_pending"


# ---------------------------------------- holdings depth (Phase 4), the rest


def test_suggested_dividends_are_added_as_one_undoable_batch(monkeypatch):
    rev = inv.add(BUY, 0)["rev"]
    with pytest.raises(inv.BookError):
        inv.add_dividends([{**BUY, "type": "buy"}], rev)  # dividends only
    divs = [{"type": "dividend", "symbol": "AAPL", "date": d, "amount": 2.5, "tax": 0, "ccy": "USD"}
            for d in ("2025-05-12", "2025-08-11")]  # fmt: skip
    b = inv.add_dividends(divs, rev)
    assert b["undo_label"] == "Added 2 dividends from Yahoo"
    raw = _raw()
    added = [e for e in raw["entries"] if e["type"] == "dividend"]
    assert {e["source"] for e in added} == {"suggested"} and len({e["batch"] for e in added}) == 1
    assert len(inv.undo(b["rev"])["entries"]) == 1


def test_the_export_has_positions_and_ledger_sheets(srv):
    from openpyxl import load_workbook

    _seed()
    with urllib.request.urlopen(srv + "/api/export-xlsx?ccy=USD") as r:
        wb = load_workbook(io.BytesIO(r.read()))
    assert wb.sheetnames[:4] == ["Read me", "Overview", "Positions", "Ledger"]
    pos = wb["Positions"]
    names = [pos.cell(row=k, column=1).value for k in range(5, pos.max_row + 1)]
    assert "Cash" in names and "Total" in names and "NVIDIA" in " ".join(n or "" for n in names)
    led = wb["Ledger"]
    assert led.cell(row=4, column=1).value == "Date"
    assert led.max_row - 4 == len(_raw()["entries"])
    dates = [led.cell(row=k, column=1).value for k in range(5, led.max_row + 1)]
    assert dates == sorted(dates)  # oldest first


# ------------------------------------------- adversarial review (Phase 4)


def test_a_ticker_yahoo_doesnt_carry_counts_as_entered_not_pending(monkeypatch):
    # No prices and no split history for it, while Yahoo answers for the
    # benchmark: a delisted line. Its trades count as entered, with a note,
    # and it never sits "pending" (which would also block its edits).
    QUOTE["OLDCO"] = "USD"
    monkeypatch.setitem(QUOTE, "OLDCO", "USD")
    _yahoo_down_for(monkeypatch, "OLDCO")
    b = inv.add({"type": "buy", "symbol": "OLDCO", "date": "2025-01-02", "qty": 10, "price": 20,
                 "qty_basis": "trade"}, 0)  # fmt: skip
    assert not b["splits_pending"] and "broken" not in b
    assert any("Yahoo has no history for" in w for w in b["warnings"])
    b = inv.add({"type": "sell", "symbol": "OLDCO", "date": "2025-02-03", "qty": 4, "price": 21,
                 "qty_basis": "trade"}, b["rev"])  # fmt: skip
    assert b["rev"] == 2


def test_in_an_outage_even_the_benchmark_is_missing_so_it_stays_pending(monkeypatch):
    _seed()
    _yahoo_down_for(monkeypatch, "NVDA")
    monkeypatch.setattr(inv, "_bench_closes", lambda period, ticker: None)
    b = inv.book()
    assert [p["symbol"] for p in b["splits_pending"]] == ["NVDA"]


def test_split_cache_survives_threads():
    import concurrent.futures as cf

    syms = [f"S{i}" for i in range(60)]
    with cf.ThreadPoolExecutor(8) as ex:
        list(ex.map(lambda k: inv.market_splits(syms[k % 7 :]), range(40)))
    saved = json.loads(paths.state_file("splits").read_text(encoding="utf-8"))
    assert set(saved) == set(syms)


def test_suggested_dividends_may_exceed_one_imports_limit():
    rev = inv.add(BUY, 0)["rev"]
    divs = [{"type": "dividend", "symbol": "AAPL", "date": "2025-02-03", "amount": 0.01,
             "ccy": "USD"}] * (inv.IMPORT_MAX + 1)  # fmt: skip
    assert (
        inv.add_dividends(divs, rev)["undo_label"]
        == f"Added {inv.IMPORT_MAX + 1} dividends from Yahoo"
    )


def test_the_export_never_writes_a_formula(srv):
    from openpyxl import load_workbook

    rev = inv.add(BUY, 0)["rev"]
    inv.add({"type": "deposit", "date": "2025-01-02", "amount": 5, "ccy": "USD",
             "note": '=HYPERLINK("http://example.com","x")'}, rev)  # fmt: skip
    with urllib.request.urlopen(srv + "/api/export-xlsx?ccy=USD") as r:
        wb = load_workbook(io.BytesIO(r.read()))
    led = wb["Ledger"]
    notes = [led.cell(row=k, column=15) for k in range(5, led.max_row + 1)]
    hit = [c for c in notes if c.value and str(c.value).startswith("=")]
    assert hit and all(c.data_type == "s" for c in hit)


def test_yield_on_cost_is_none_without_dividends_and_contrib_is_never_nan():
    assert inv._yield_on_cost("AAPL", "USD", 10, 1000.0, {"USD": np.ones(3)}) is None
    daily = pd.DataFrame({"A": [0.0, -1.0], inv.CASH_BUCKET: [0.0, 0.0]})
    out = inv._contrib_linked(daily, -1)
    assert all(v is None or np.isfinite(v) for v in out.values())
    json.dumps(out, allow_nan=False)


def test_analyst_upside_comes_from_any_saved_portfolio_when_the_book_tab_hasnt_loaded():
    from convexity import persistence

    persistence.save_view(
        "Tech", "AAPL", [{"symbol": "AAPL", "price": 200.0, "target_mean_price": 250.0}]
    )
    b = inv.add(BUY, 0)
    assert b["positions"][0]["upside_pct"] == pytest.approx(25.0)


# ------------------------------------------------- the page, fast (2026-10-12)


def test_page_computes_once_for_concurrent_callers(monkeypatch):
    inv.add(BUY, 0)
    inv._PAGE_MEMO.clear()
    calls, real = [], inv.book

    def slow(ccy=None):
        calls.append(ccy)
        time.sleep(0.2)
        return real(ccy)

    monkeypatch.setattr(inv, "book", slow)
    out = []
    ts = [threading.Thread(target=lambda: out.append(inv.page("USD"))) for _ in range(3)]
    for t in ts:
        t.start()
    for t in ts:
        t.join()
    # One computation shared (each answer is a copy with its own Attention rail).
    assert (
        len(calls) == 1
        and len(out) == 3
        and all(o["positions"] is out[0]["positions"] for o in out)
    )
    # A change bumps the rev: the next page is computed again.
    inv.add(BUY | {"symbol": "MSFT"}, 1)
    calls.clear()
    assert inv.page("USD")["rev"] == 2 and calls == []  # the change already computed it
    assert inv.page("USD", fresh=True)["rev"] == 2 and len(calls) == 1


def test_the_last_good_page_is_saved_and_only_served_for_the_same_book():
    assert inv.cached_page("USD") is None  # nothing saved yet
    b = inv.add(BUY, 0)
    c = inv.cached_page("USD")
    assert c["stale"] is True and c["computed_at"] and c["rev"] == b["rev"] == 1
    assert c["headline"]["value"] == pytest.approx(b["headline"]["value"])
    assert inv.cached_page("GBP") is None  # another currency: not this page
    inv.add(BUY | {"symbol": "MSFT"}, 1)
    path = paths.state_file("book_cache")
    body = json.loads(path.read_text(encoding="utf-8"))
    body["rev"] = 1  # an older book's page
    path.write_text(json.dumps(body), encoding="utf-8")
    assert inv.cached_page("USD") is None
    path.write_text("{not json", encoding="utf-8")
    assert inv.cached_page("USD") is None  # a damaged cache is ignored


def test_an_unhealthy_page_is_never_saved(monkeypatch):
    inv.add(BUY | {"symbol": "NVDA", "qty_basis": "trade"}, 0)
    inv.add(BUY, 1)
    paths.state_file("book_cache").unlink()
    monkeypatch.setattr(inv, "_fetch_splits", lambda s: None)  # Yahoo down, nothing known
    inv._SPLITS.clear()
    inv._SPLITS_GOOD.clear()
    paths.state_file("splits").unlink(missing_ok=True)
    monkeypatch.setattr(inv, "_splits_loaded", None)
    b = inv.page("USD", fresh=True)
    assert b["splits_pending"] and inv.cached_page("USD") is None


def test_the_cached_route_answers_at_once(srv):
    with urllib.request.urlopen(srv + "/api/investments?cached=1&ccy=USD") as r:
        assert json.loads(r.read()) == {"nocache": True}
    _post(srv, "/api/investments/entries", {"base_rev": 0, "entry": BUY})
    with urllib.request.urlopen(srv + "/api/investments?cached=1&ccy=USD") as r:
        c = json.loads(r.read())
    assert c["stale"] is True and c["rev"] == 1


def test_warm_skips_an_empty_book(monkeypatch):
    monkeypatch.setattr(inv, "page", lambda ccy=None: pytest.fail("no book, no work"))
    inv.warm("EUR")


# ------------------------------------------------------- attention (Phase 7)


def test_page_carries_xray_and_the_rail(srv):
    _seed()
    b = inv.page("USD")
    held = [p for p in b["positions"] if not p.get("pending")]
    assert held and all("xray" in p for p in held)
    att = b["attention"]
    assert {"high", "info", "upcoming", "unread", "rules"} <= set(att)
    # The memo and the cached page never hold the rail (it follows the
    # thresholds at the time of asking).
    assert "attention" not in next(iter(inv._PAGE_MEMO.values()))[1]
    cached = json.loads(paths.state_file("book_cache").read_text(encoding="utf-8"))["book"]
    assert "attention" not in cached
    assert "attention" in inv.cached_page("USD")
    assert "attention" in inv.page("USD")  # the memo's answer too


def test_attention_route_never_computes(srv, monkeypatch):
    _seed()
    calls = []
    monkeypatch.setattr(inv, "book", lambda ccy=None: calls.append(ccy) or {"empty": True})
    with urllib.request.urlopen(srv + "/api/investments/attention?ccy=USD") as r:
        # A book with no page computed yet: "pending", so the badge asks again;
        # the rules still come (Settings works before the first page).
        a = json.loads(r.read())
        assert a["pending"] is True and a["unread"] == 0 and a["high"] == [] and a["rules"]
    assert calls == []


def test_attention_routes(srv):
    _seed()
    inv.page("USD")
    rev = _raw()["rev"]
    code, att = _post(srv, "/api/investments/attention/rules", {"rules": {"holding": {"value": 1}}})
    assert code == 200
    hold = [it for it in att["high"] if it["rule"] == "holding"]
    assert hold, "every holding is above 1%"
    code, att2 = _post(srv, "/api/investments/attention/dismiss",
                       {"key": hold[0]["key"], "level": hold[0]["level"]})  # fmt: skip
    assert code == 200 and hold[0]["key"] not in [it["key"] for it in att2["high"]]
    code, att3 = _post(srv, "/api/investments/attention/seen",
                       {"items": [[it["key"], it["level"]] for it in att2["high"] + att2["info"]]})  # fmt: skip
    assert code == 200 and att3["unread"] == 0
    for path, body in (("rules", {"rules": {"nope": {}}}), ("dismiss", {"key": "", "level": 1}),
                       ("seen", {"items": "x"}), ("other", {})):  # fmt: skip
        code, err = _post(srv, "/api/investments/attention/" + path, body)
        assert code == 400, path
    code, _ = _post(srv, "/api/investments/attention/rules", {"rules": {"holding": {"value": 1}}},
                    ctype="text/plain")  # fmt: skip
    assert code == 403
    assert _raw()["rev"] == rev  # the rail's preferences never touch the book


def test_plan_routes_answer_the_rail_and_never_touch_the_book(srv):
    _seed()
    page = inv.page("USD")
    rev, x = _raw()["rev"], page["positions"][0]
    sym, price = x["symbol"], x["price"]
    code, att = _post(srv, "/api/investments/attention/plan",
                      {"symbol": sym, "plan": {"target": price * 0.9, "why": "demo"}, "price": price})  # fmt: skip
    assert code == 200 and att["plans"][sym]["why"] == "demo"
    tgt = [it for it in att["high"] if it["rule"] == "target"]
    assert [it["symbols"] for it in tgt] == [[sym]]
    code, att = _post(
        srv, "/api/investments/attention/snooze", {"key": tgt[0]["key"], "until": "2099-01-01"}
    )
    assert code == 400  # too far
    until = (date.today() + timedelta(days=1)).isoformat()
    code, att = _post(
        srv, "/api/investments/attention/snooze", {"key": tgt[0]["key"], "until": until}
    )
    assert code == 200 and not [it for it in att["high"] if it["rule"] == "target"]
    code, att = _post(
        srv, "/api/investments/attention/alert", {"symbol": sym, "level": price * 2, "price": price}
    )
    assert code == 200 and att["plans"][sym]["alerts"][0]["side"] == "above"
    aid = att["plans"][sym]["alerts"][0]["id"]
    code, att = _post(srv, "/api/investments/attention/alert/remove", {"symbol": sym, "id": aid})
    assert code == 200 and "alerts" not in att["plans"][sym]
    code, att = _post(srv, "/api/investments/attention/plan/clear", {"symbol": sym})
    assert code == 200 and sym not in att["plans"]
    for path, body in (("plan", {"symbol": "ZZZZ", "plan": {"target": 1}, "price": 1}),
                       ("plan", {"symbol": sym, "plan": "x", "price": price}),
                       ("alert", {"symbol": sym, "level": "high", "price": price}),
                       ("snooze", {"key": "k", "until": "tomorrow"})):  # fmt: skip
        code, _ = _post(srv, "/api/investments/attention/" + path, body)
        assert code == 400, (path, body)
    paths.state_file("plans").write_text("{oops", encoding="utf-8")
    code, err = _post(
        srv, "/api/investments/attention/plan", {"symbol": sym, "plan": {"target": 1}, "price": 1}
    )
    assert code == 409 and paths.state_file("plans").read_text(encoding="utf-8") == "{oops"
    assert _raw()["rev"] == rev  # plans live beside the book, never in it


def test_plan_prices_come_from_the_server_and_sold_out_plans_clear(srv, monkeypatch):
    from convexity import attention

    _seed()
    page = inv.page("USD")
    x = page["positions"][0]
    sym, price = x["symbol"], x["price"]
    # A stale page sends a price far off: the side is fixed against the server's.
    code, att = _post(srv, "/api/investments/attention/alert",
                      {"symbol": sym, "level": price * 1.01, "price": price * 2})  # fmt: skip
    assert code == 200 and att["plans"][sym]["alerts"][0]["side"] == "above"
    code, att = _post(srv, "/api/investments/attention/plan",
                      {"symbol": sym, "plan": {"review": price / 2}, "price": 1})  # fmt: skip
    assert code == 200 and att["plans"][sym]["set_price"] == pytest.approx(price)
    monkeypatch.setattr(inv, "book_symbols", list)  # sold out
    aid = attention.load_plans()[sym]["alerts"][0]["id"]
    assert (
        _post(srv, "/api/investments/attention/alert/remove", {"symbol": sym, "id": aid})[0] == 200
    )
    assert _post(srv, "/api/investments/attention/plan/clear", {"symbol": sym})[0] == 200
    assert sym not in attention.load_plans()
    assert _post(srv, "/api/investments/attention/plan",
                 {"symbol": sym, "plan": {"target": 1}, "price": 1})[0] == 400  # fmt: skip


def test_refresh_also_fetches_event_dates(monkeypatch):
    from convexity import attention

    _seed()
    seen = []
    monkeypatch.setattr(
        attention, "refresh_events", lambda syms, **k: seen.append((sorted(syms), k))
    )
    syms = inv.refresh_symbols()
    inv.refresh_market(syms)
    bench = inv.INV_BENCHMARKS[inv.bench_key()][0]
    assert seen and seen[0][1]["force"] is True and bench not in seen[0][0]
    assert set(seen[0][0]) == set(syms) - {bench}
